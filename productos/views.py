
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse_lazy
from django.views.generic import (
    CreateView,
    DeleteView,
    DetailView,
    ListView,
    UpdateView,
)

from .forms import ProductoForm
from .models import Producto


# ── VISTAS DE PRODUCTOS ────────────────────────────────────────────────

class ProductoListView(LoginRequiredMixin, ListView):
    model = Producto
    template_name = 'productos/lista.html'
    context_object_name = 'productos'
    paginate_by = 10

    def get_queryset(self):
        qs = super().get_queryset().select_related(
            'categoria',
            'proveedor',
        )

        q = self.request.GET.get('q')
        cat = self.request.GET.get('categoria')

        if q:
            qs = qs.filter(nombre__icontains=q)

        if cat:
            qs = qs.filter(categoria_id=cat)

        return qs


class ProductoDetailView(LoginRequiredMixin, DetailView):
    model = Producto
    template_name = 'productos/detalle.html'
    context_object_name = 'producto'
    pk_url_kwarg = 'producto_id'


class ProductoCreateView(LoginRequiredMixin, CreateView):
    model = Producto
    form_class = ProductoForm
    template_name = 'productos/crear.html'
    success_url = reverse_lazy('productos:lista')


class ProductoUpdateView(LoginRequiredMixin, UpdateView):
    model = Producto
    form_class = ProductoForm
    template_name = 'productos/editar.html'
    pk_url_kwarg = 'producto_id'

    def get_success_url(self):
        return reverse_lazy(
            'productos:detalle',
            kwargs={'producto_id': self.object.pk},
        )


class ProductoDeleteView(LoginRequiredMixin, DeleteView):
    model = Producto
    template_name = 'productos/eliminar.html'
    pk_url_kwarg = 'producto_id'
    success_url = reverse_lazy('productos:lista')


# ── VISTA W12: EXPORTACIÓN EXCEL ───────────────────────────────────────

@login_required
def exportar_productos_excel(request):
    """Exporta el inventario de productos a un archivo Excel (.xlsx)."""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    # Encabezados con formato AzulERP
    headers = [
        "ID",
        "Nombre",
        "Categoría",
        "Proveedor",
        "Precio ($)",
        "Stock",
        "Estado",
    ]
    ws.append(headers)

    fill = PatternFill(
        start_color="0A2342",
        end_color="0A2342",
        fill_type="solid",
    )
    font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")

    # Datos de productos
    productos = Producto.objects.select_related(
        'categoria',
        'proveedor',
    ).all()

    for p in productos:
        ws.append([
            p.pk,
            p.nombre,
            p.categoria.nombre if p.categoria else "—",
            p.proveedor.nombre if p.proveedor else "—",
            float(p.precio),
            p.stock,
            "Activo" if p.activo else "Inactivo",
        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )
    response['Content-Disposition'] = (
        'attachment; filename="inventario_productos.xlsx"'
    )

    wb.save(response)
    return response
